# PDP 모수 1단계 거르기: 상품명·주소만으로 할 수 있는 것(브랜드, 제품군, 리리부 단어, 같은 주소 중복).
# 페이지를 열어야 하는 판단(마켓플레이스, 실제 도착 주소)은 '페이지 확인 필요'로 표시만 한다.
import csv, glob, html, json, re, sys, collections
from urllib.parse import urlparse
import openpyxl

ROOT = '/Users/codestepseunghee/codestep/lg-online-retailer-analysis'
PDP = ROOT + '/pdp 모수 추출'
OUT = sys.argv[1]

# ---------- 기준 자료 ----------
wb = openpyxl.load_workbook(glob.glob(ROOT + '/관련문서/*v2.2*.xlsx')[0], data_only=True)

COUNTRY_ALIAS = {'SaudiArabia': 'Saudi Arabia', 'UK': 'United Kingdom', 'Czech': 'Czech Republic'}
def cnorm(c):
    c = (c or '').strip()
    return COUNTRY_ALIAS.get(c, c)

# 경쟁브랜드 (국가, REF|W/M|DW) -> 브랜드 집합
BRANDS = collections.defaultdict(set)
for row in wb['경쟁브랜드'].iter_rows(min_row=5, values_only=True):
    vals = [v for v in row if v not in (None, '')]
    if len(vals) == 4:
        _, country, cat, brand = vals
        BRANDS[(cnorm(country), cat.strip())].add(brand.strip().upper())
# 독일·영국은 두레이 #110 (D062)
for cat, bl in {'W/M': ['LG', 'SAMSUNG', 'BOSCH', 'SIEMENS'], 'REF': ['LG', 'SAMSUNG', 'BOSCH', 'LIEBHERR']}.items():
    BRANDS[('Germany', cat)] = set(bl)
for cat in ['W/M', 'REF']:
    BRANDS[('United Kingdom', cat)] = {'LG', 'SAMSUNG', 'BOSCH', 'HISENSE'}

# 제외키워드(리리부): 오른쪽 25개국 표만 쓴다(왼쪽 7개국 표를 모두 포함)
KW = collections.defaultdict(list)
ws = wb['제외키워드(리리부)']
hdr = None
for row in ws.iter_rows(values_only=True):
    row = list(row)
    if hdr is None:
        idx = [i for i, v in enumerate(row) if v == 'CNTRY']
        if len(idx) == 2:
            hdr = idx[1]
        continue
    c, typ, kw = row[hdr], row[hdr + 1], row[hdr + 2]
    if c and typ and kw:
        KW[cnorm(c)].append((typ, html.unescape(str(kw)).strip().lower()))

def kw_hit(country, text):
    t = text.lower()
    for typ, kw in KW.get(country, []):
        # 단어 앞부분 경계만 본다(odnowion·repasovan 같은 어간 단어 때문)
        if re.search(r'(?<![\w])' + re.escape(kw), t):
            return typ, kw
    return None

# ---------- 브랜드 찾기 (검색 흐름) ----------
def brand_pattern(b):
    b = b.lower().replace('&', ' ')
    parts = [re.escape(p) for p in b.split()]
    return re.compile(r'(?<![a-z0-9])' + r'[\s\-_&+]*'.join(parts) + r'(?![a-z0-9])')
BRAND_RE = {}
def brands_in(text, cand):
    t = text.lower()
    found = []
    for b in cand:
        if b not in BRAND_RE:
            BRAND_RE[b] = brand_pattern(b)
        if BRAND_RE[b].search(t):
            found.append(b)
        elif b.lower() in AR_BRAND and re.search(AR_BRAND[b.lower()], ' '.join(t.split()[:3])):
            # 아랍어 상품명은 브랜드가 맨 앞에 온다. 앞 3단어만 봐서 '엘지 모터 탑재' 같은 문구는 안 잡는다
            found.append(b)
    return found
# 아랍어 브랜드 표기 (이집트·모로코·UAE 상품명용, AI가 추가)
AR_BRAND = {'lg': r'(?:^|\s)[اإ]ل ?جي(?:\s|$)', 'samsung': r'سامسون[جغ]', 'toshiba': r'توشيبا', 'sharp': r'شارب',
            'zanussi': r'زانوسي', 'beko': r'بيكو', 'siera': r'سييرا|سيرا', 'whirlpool': r'ويرلبول|ويرلپول',
            'bosch': r'بوش', 'hisense': r'هايسنس', 'haier': r'هاير', 'midea': r'ميديا', 'ariston': r'اريستون'}

def slug(url):
    p = urlparse(url)
    return p.path.replace('/', ' ').replace('-', ' ').replace('_', ' ')

# ---------- 제품군 단어 (검색 흐름, AI가 정한 목록) ----------
NEG = [  # 대상 제품군 아님 -> 제외
    ('식기세척기', r'dishwasher|vaatwas|myčk|zmywar|mosogatógép|lavar loiça|lavavajilla|lave[\s-]?vaisselle|^lv |غسالة (?:ال)?(?:اطباق|أطباق|صحون)'),
    ('헤어드라이어', r'hair ?dryer|secador(?:a)? de (?:cabello|pelo)|vysoušeč vlasů|suszarka do włosów|hajszárító|سشوار|مجفف شعر|sèche[\s-]?cheveux|secador profesional|secador ionico'),
    ('빨래건조대', r'wieszak|sušák|suszarka na (?:pranie|ubrania)|tendedero|secador de ropa plegable|ruhaszárító állvány|stendibiancheria'),
    ('스타일러', r'styler'),
    ('기타 가전', r'televis|\btv\b|microwave|mikrovln|mikrofal|micro[\s-]?ondas|hob\b|plaque|varná deska|płyta|expresso|espresso|kávovar|écouteur|earbuds|vacuum|vysavač|odkurzacz|aspirador|klimatizace|klimatyzac|air condition|cooker hood|okap|digestoř|oven\b|trouba|piekarnik'),
]
POS = {
    'REF': r'fridge|refrigerat|réfrigérat|refrigerador|koelkast|koel[\s-]?vries|chladni|lednic|lodówk|hűtő|frigor|frigo\b|ثلاج|side by side|french door|top mount|bottom mount|\bcomb\b|\bref\b|americain|américain|multi[\s-]?portes|\bsbs\b|mini ?bar|cooler',
    'W/M': r'washer|washing|wasmachine|wasautomaat|pračk|pralk|mosógép|mosó-szárító|máquina de lavar(?! loiça)|lavadora|lavasecadora|غسال|lave[\s-]?linge|machine à laver|\bmal\b|twin tub|semiautom',
    'DRY': r'dryer|droger|sušičk|suszark|szárító|máquina de secar|secadora|secarropa|نشاف|مجفف|sèche[\s-]?linge|\bsl\b|\bsech',
}
FREEZER_HEAD = r'^(?:\S+\s+){0,2}(?:freezer|vriezer|mrazák|mraznička|zamrażarka|fagyasztószekrény|fagyasztóláda|congelador|congélateur|congelateur|chest freezer|upright freezer|deep freezer|ديب فريزر|فريزر)'
FREEZER = r'freezer|vriezer|mrazák|mraznička(?! s)|zamrażark|fagyasztó(?!s)|congelador|congelateur|congélateur|ديب ?فريزر|فريزر'
# 시트 밖 부품 의심 단어(AI 추가) -> 제외하지 않고 '확인 필요'
PART_SUSPECT = r'filtr|filter|tampa|cover|hose|wąż|hadice|tömlő|mangueira|manguera|bracket|soporte|podstav|podest|pedestal|stacking|łącznik|spojovac|gasket|uszczelk|těsnění|półk|polic|shelf|tray|cleaner|limpeza|odkamieniacz|odvápň|tablet|kapsuł|capsul|deterg|proszek|płyn|liquid|líquido|sacos?|worki|bag|remote|pilot|cable|kabel|replacement|náhradn|zamienn|alkatrész|repuesto|peça|diodo|lâmpada|bulb|kit(?!\w)|adapter|adaptér|adapte|organizer|cesto|gaveta|basket|kosz|koš|ice maker|kostkarka|water ?filter|حامل|فلتر|قاعدة'
def head(name, n=4):
    # 상품 종류는 상품명 앞쪽에 오므로 앞 n단어만 본다
    return ' '.join(re.split(r'\s+', name.strip().lower())[:n])

def search_cat(fcol):
    f = fcol.lower()
    if 'refrig' in f: return 'REF'
    if 'wash' in f: return 'W/M'
    if 'dryer' in f: return 'DRY'
    return '?'

def category_check(cat, text):
    t = text.lower()
    for name, pat in NEG:
        if re.search(pat, t):
            return '제외', name
    if re.search(FREEZER_HEAD, head(t)) and not re.search(r'combi|kombin|s mrazák|s mrazničk|z zamraż|fagyasztós|réfrigérateur|refrigera|lodówk|chladni|hűtő|geladeira|ثلاج|fridge|koel|frigor', t):
        return '통과', '냉동고'
    poss = [k for k, p in POS.items() if re.search(p, t)]
    if poss:
        return '통과', '/'.join(poss)
    return '확인', '제품군 단어 없음'

# ---------- 입력 읽기 ----------
rows = []
def osm_rows(fname, keep):
    for x in csv.DictReader(open(PDP + '/' + fname, encoding='utf-8-sig')):
        if not keep(x):
            continue
        url = x.get('CRAWL_URL') if x.get('CRAWL_URL') not in (None, '', '\\N') else x['RAW_URL']
        rows.append({
            '흐름': 'OSM', '출처파일': fname, '국가': cnorm(x['CNTRY']), '리테일러': x['RTLR_NAME'],
            '제품군': x['DIV'], 'OSM브랜드': x['BRAND'], '상품명': x['PRODUCT_DESC'], 'URL': url,
            '원본URL': x['RAW_URL'], '기준일': x['RPT_DATE'], '검색어': '', '엑셀행번호': '', '광고여부': '',
            'OSM_RK': x['RK'], 'OSM_SOURCE_TYPE': x['SOURCE_TYPE'],
        })
osm_rows('OSM_유럽13개리테일러_PDP후보_8160건_20261026.csv', lambda x: x['CNTRY'] not in ('Germany', 'UK'))
osm_rows('OSM_독일영국5개리테일러_PDP후보_1889건_20261029.csv', lambda x: True)
osm_rows('OSM_추가9개리테일러_PDP후보_4942건_20261028.csv', lambda x: True)
SF = '검색결과_20261028_734afcdd/상품카드.csv'
for x in csv.DictReader(open(PDP + '/' + SF, encoding='utf-8-sig')):
    rows.append({
        '흐름': '검색', '출처파일': SF, '국가': cnorm(x['국가']), '리테일러': x['리테일러'].strip(),
        '제품군': search_cat(x['F열제품군']), 'OSM브랜드': '', '상품명': x['상품명'], 'URL': x['카드PDP링크'],
        '원본URL': x['카드PDP링크'], '기준일': x['수집시각'][:10], '검색어': x['J열검색어'], '엑셀행번호': x['엑셀행번호'],
        '광고여부': x['광고여부'], 'OSM_RK': '', 'OSM_SOURCE_TYPE': '',
    })

FR58 = set()
for x in csv.DictReader(open(OUT + '/france_58_74.csv', encoding='utf-8-sig')):
    FR58.add(x['URL'])
    rows.append({
        '흐름': '#58', '출처파일': '두레이 #58 본문 표', '국가': 'France', '리테일러': x['Retailer'].upper(),
        '제품군': {'WM': 'W/M'}.get(x['Category'], x['Category']), 'OSM브랜드': x['Brand'], '상품명': x['Product Name'],
        'URL': x['URL'], '원본URL': x['URL'], '기준일': '2026-08', '검색어': '', '엑셀행번호': '', '광고여부': '',
        'OSM_RK': '', 'OSM_SOURCE_TYPE': '',
    })

# 페이지를 열어야 판단하는 리테일러별 규칙(README 6장). 규칙이 문서로 있는 곳만 적는다.
PAGE_RULE = {
    ('Spain', 'MEDIAMARKT'): '"Vendido y enviado por" 있으면 제외(#74)',
    ('Spain', 'ELCORTEINGLES_ELECTRONICA'): '"Vendido por <외부 판매자>"면 제외(#94)',
    ('Italy', 'MEDIAWORLD'): '왼쪽 위 LG 링크 여부(#72)',
    ('France', 'BOULANGER'): '"vendu & expédié par Boulanger" 아니면 제외(#58)',
    ('France', 'DARTY'): 'Darty 판매만(#58 첨부)',
}

def cat_key(country, div):
    if div in ('REF',): return 'REF'
    if div in ('W/M', 'DRY'): return 'W/M'   # 건조기는 세탁 브랜드 목록을 쓴다(AI 해석)
    if div in ('D/W',): return 'DW'
    return div

# ---------- 판정 ----------
def judge(r, france_mode):
    """france_mode: 'A'(#58 74개) | 'B'(v2.2 경쟁브랜드). 반환: (상태, 단계, 사유, 메모)"""
    c, ret, div, name, url = r['국가'], r['리테일러'], r['제품군'], r['상품명'] or '', r['URL']
    memo = []
    note = []  # 판정에 영향 없는 참고 메모
    # 프랑스 A: #58 목록에 있는 것만
    if c == 'France' and france_mode == 'A':
        if r['흐름'] == 'OSM':
            base = url.split('?')[0]
            why = '#58 목록 주소에 ?ofmp 붙은 형태(목록 주소 행이 따로 있음)' if base in FR58 and base != url else '#58 8월 모수 74개에 없음'
            return '제외', '①브랜드', '프랑스 A: ' + why if url not in FR58 else '프랑스 A: #58 목록 행으로 대체', ''
    if r['흐름'] == '#58' and france_mode == 'B':
        return '제외', '①브랜드', '프랑스 B에서는 #58 목록 행을 쓰지 않음', ''
    # ① 브랜드
    ck = cat_key(c, div)
    target = BRANDS.get((c, ck), set())
    if r['흐름'] in ('OSM', '#58'):
        b = (r['OSM브랜드'] or '').strip().upper()
        if div == 'D/W' and c != 'Spain':
            pass  # ②에서 처리
        elif not target:
            return '확인', '①브랜드', f'경쟁브랜드 목록 없음({c} {ck})', ''
        elif b not in target:
            return '제외', '①브랜드', f'대상 브랜드 아님({b})', ''
    else:
        if div == '?':
            return '확인', '②제품군', '검색어 제품군 해석 불가', ''
        text = name + ' ' + slug(url)
        found = brands_in(text, target)
        if not found:
            return '제외', '①브랜드', '대상 브랜드 이름 없음(상품명·주소)', ''
        if len(found) > 1:
            memo.append('대상 브랜드 여러 개: ' + ','.join(found))
    # ② 제품군
    if div == 'D/W' and c != 'Spain':
        return '제외', '②제품군', '식기세척기는 스페인만(D061)', ''
    if r['흐름'] == '검색':
        st, why = category_check(div, name + ' ' + slug(url))
        if st == '제외':
            return '제외', '②제품군', f'대상 제품군 아님({why})', ''
        if st == '통과' and why == '냉동고':
            note.append('단독 냉동고(냉장고에 포함, 10-29 사용자 결정)')
        if st == '확인' and why != '제품군 단어 없음':
            memo.append(why)
        elif why == '제품군 단어 없음':
            note.append('제품군 단어 없음(검색어 제품군 따름)')
    else:
        if re.search(FREEZER_HEAD, head(name)) and not re.search(r'combi|kombin|frigo|réfrigérateur|refrigera|lodówk|chladni|hűtő|fridge|kühl|koel|kyl|ثلاج|geladeira|冰箱|tủ lạnh', name.lower()):
            note.append('단독 냉동고(냉장고에 포함, 10-29 사용자 결정)')
    # ③ 리리부 (시트 단어)
    hit = kw_hit(c, name)
    if hit:
        return '제외', '③리리부', f'{hit[0]} 단어 "{hit[1]}"(제외키워드 시트)', ''
    if c == 'Sweden' and 'outlet' in url.lower():
        return '제외', '③리리부', 'URL에 outlet(D046)', ''
    if c == 'Canada' and ('langId=-25' in url):
        return '제외', '③리리부', '불어 페이지(D050)', ''
    if re.search(r'(?<![\w])(?:' + PART_SUSPECT + ')', head(name)):
        memo.append('부품 의심 단어(시트 밖, AI 추가)')
    if ret == 'DARTY' and 'ofmp=' in url:
        note.append('주소에 ofmp(외부 판매자 상품 주소로 보임, 페이지 확인 필요)')
    return ('확인' if memo else '통과'), '', '', '; '.join(memo + note)

def run(france_mode):
    out = []
    for i, r in enumerate(rows, 1):
        st, step, why, memo = judge(r, france_mode)
        o = dict(r); o['행번호'] = i
        o['판정'] = {'통과': '남김', '제외': '제외', '확인': '확인 필요'}[st]
        o['제외단계'] = step; o['제외사유'] = why; o['메모'] = memo
        o['페이지확인규칙'] = PAGE_RULE.get((r['국가'], r['리테일러']), '')
        out.append(o)
    # ④ 같은 주소 중복: 제외 아닌 행끼리, 같은 리테일러 안에서 완전히 같은 URL
    first = {}
    for o in out:
        if o['판정'] == '제외':
            continue
        k = (o['국가'], o['리테일러'], o['URL'])
        if k in first:
            o['판정'] = '제외'; o['제외단계'] = '④중복'; o['제외사유'] = f'같은 주소(대표 행 {first[k]})'
        else:
            first[k] = o['행번호']
    return out

COLS = ['행번호', '흐름', '국가', '리테일러', '제품군', 'OSM브랜드', '상품명', 'URL', '원본URL', '기준일', '검색어', '엑셀행번호',
        '광고여부', 'OSM_RK', 'OSM_SOURCE_TYPE', '출처파일', '판정', '제외단계', '제외사유', '메모', '페이지확인규칙']
summary = {}
for mode in ['A', 'B']:
    res = run(mode)
    with open(f'{OUT}/step1_프랑스{mode}.csv', 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, COLS); w.writeheader(); w.writerows(res)
    summary[mode] = res
json.dump({'kw_countries': sorted(KW), 'brand_keys': sorted(f'{a}|{b}' for a, b in BRANDS)}, open(OUT + '/step1_meta.json', 'w'), ensure_ascii=False, indent=1)

# 리테일러별 집계(프랑스만 A/B 따로)
agg = collections.OrderedDict()
for mode, res in summary.items():
    for o in res:
        if mode == 'B' and o['국가'] != 'France':
            continue
        key = (o['흐름'], o['국가'], o['리테일러'] + (f' [{mode}]' if o['국가'] == 'France' else ''))
        a = agg.setdefault(key, collections.Counter())
        a['전체'] += 1
        a[o['판정']] += 1
        if o['판정'] == '제외':
            a[o['제외단계']] += 1
with open(OUT + '/step1_집계.csv', 'w', encoding='utf-8-sig', newline='') as f:
    cols = ['흐름', '국가', '리테일러', '전체', '①브랜드', '②제품군', '③리리부', '④중복', '남김', '확인 필요']
    w = csv.writer(f); w.writerow(cols)
    for (fl, c, rt), a in sorted(agg.items()):
        w.writerow([fl, c, rt] + [a.get(k, 0) for k in cols[3:]])
print(open(OUT + '/step1_집계.csv', encoding='utf-8-sig').read())
